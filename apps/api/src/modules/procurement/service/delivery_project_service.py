"""Multi-site delivery projects: one tracker for a programme spread across circles/sites.

GST rules force one customer PO per destination, so a 100-server / 10-site
order arrives as 10 POs. The project groups them: every site carries its own
PO, OVF, vendor PO, milestone and dates, the whole tracker round-trips through
Excel (keyed by the Site Key column), and the customer gets a read-only link.
"""

from __future__ import annotations

import base64
import binascii
import secrets
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from uuid import UUID

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from core.exceptions import ConflictException, ForbiddenException, NotFoundException, ValidationException
from modules.foundation.domain.value_objects import TenantContext
from modules.procurement.adapters.crm_adapter import ProcurementCrmAdapter
from modules.procurement.models.delivery_project import ProcDeliveryProject, ProcDeliveryProjectSite
from modules.procurement.models.order import ProcOrderHeader
from modules.procurement.service.procurement_scope_validator import ProcurementScopeValidator
from modules.procurement.service.scm_delivery_notification_service import DELIVERY_MILESTONES

MILESTONE_LABELS: dict[str, str] = {
    "order_placed": "Order placed",
    "ready_at_factory": "Ready at factory",
    "dispatched_from_factory": "Dispatched from factory",
    "received_in_india": "Received in India",
    "received_at_warehouse": "Received at warehouse",
    "dispatched_to_site": "Dispatched to site",
    "reached_site": "Reached site",
    "installed": "Installed",
}
SITE_STATUSES = ("pending", "in_progress", "delivered", "installed", "on_hold", "cancelled")
PROJECT_STATUSES = ("active", "on_hold", "completed")

EXCEL_COLUMNS: tuple[tuple[str, str], ...] = (
    ("site_key", "Site Key (do not edit)"),
    ("circle", "Circle"),
    ("site_code", "Site Code"),
    ("site_name", "Site Name"),
    ("address", "Address"),
    ("state", "State"),
    ("gstin", "GSTIN"),
    ("customer_po_number", "Customer PO"),
    ("item_summary", "Items"),
    ("quantity", "Qty"),
    ("milestone", "Milestone"),
    ("status", "Status"),
    ("expected_delivery_date", "Expected Delivery"),
    ("actual_delivery_date", "Actual Delivery"),
    ("awb_number", "AWB / Docket"),
    ("delay_reason", "Delay Reason"),
    ("last_note", "Note"),
)
_EDITABLE = {
    "circle", "site_code", "site_name", "address", "state", "gstin", "customer_po_number",
    "item_summary", "quantity", "milestone", "status", "expected_delivery_date",
    "actual_delivery_date", "awb_number", "delay_reason", "last_note",
}
_TYPED_FIELDS = {"quantity", "expected_delivery_date", "actual_delivery_date"}
_MAX_IMPORT_BYTES = 8 * 1024 * 1024


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _milestone_from(value: Any) -> str | None:
    text = str(value or "").strip().lower()
    if not text:
        return None
    for key, label in MILESTONE_LABELS.items():
        if text in (key, label.lower()):
            return key
    raise ValidationException(f"Unknown milestone '{value}'")


def _date_from(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValidationException(f"Could not read date '{value}' (use DD-MM-YYYY)")


def _status_for_milestone(milestone: str, current: str) -> str:
    if current in ("on_hold", "cancelled"):
        return current
    if milestone == "installed":
        return "installed"
    if milestone == "reached_site":
        return "delivered"
    if milestone == "order_placed":
        return "pending"
    return "in_progress"


class DeliveryProjectService:
    def __init__(self, db: Session) -> None:
        self._db = db
        self._scope = ProcurementScopeValidator(db)
        self._crm = ProcurementCrmAdapter(db)

    # -- projects -----------------------------------------------------------------
    def _company(self, ctx: TenantContext) -> UUID:
        cid = self._scope.resolve_company_id(ctx, None)
        if cid is None:
            raise ForbiddenException("Company context required")
        return cid

    def _get(self, ctx: TenantContext, project_id: UUID) -> ProcDeliveryProject:
        row = self._db.get(ProcDeliveryProject, project_id)
        if row is None or row.is_deleted or row.tenant_id != ctx.tenant_id:
            raise NotFoundException("Delivery project not found")
        self._scope.validate_company_access(ctx, row.company_id)
        return row

    def _sites(self, project_id: UUID) -> list[ProcDeliveryProjectSite]:
        stmt = select(ProcDeliveryProjectSite).where(
            ProcDeliveryProjectSite.project_id == project_id,
            ProcDeliveryProjectSite.is_deleted.is_(False),
        )
        return list(
            self._db.scalars(
                stmt.order_by(
                    ProcDeliveryProjectSite.circle,
                    ProcDeliveryProjectSite.sort_order,
                    ProcDeliveryProjectSite.site_name,
                )
            ).all()
        )

    @staticmethod
    def _summary(sites: list[ProcDeliveryProjectSite]) -> dict[str, Any]:
        today = date.today()
        delayed = [
            s
            for s in sites
            if s.status not in ("delivered", "installed", "cancelled")
            and s.expected_delivery_date is not None
            and s.expected_delivery_date < today
        ]
        return {
            "site_count": len(sites),
            "delivered_count": sum(1 for s in sites if s.status in ("delivered", "installed")),
            "installed_count": sum(1 for s in sites if s.status == "installed"),
            "delayed_count": len(delayed),
            "on_hold_count": sum(1 for s in sites if s.status == "on_hold"),
            "by_milestone": dict(Counter(s.milestone for s in sites)),
            "circles": sorted({s.circle for s in sites if s.circle}),
        }

    def _project_dto(self, project: ProcDeliveryProject, sites: list[ProcDeliveryProjectSite] | None = None) -> dict[str, Any]:
        rows = sites if sites is not None else self._sites(project.id)
        return {
            "id": project.id,
            "project_code": project.project_code,
            "name": project.name,
            "company_account_id": project.company_account_id,
            "customer_name": project.customer_name,
            "opportunity_id": project.opportunity_id,
            "tracking_token": project.tracking_token,
            "public_tracking_enabled": project.public_tracking_enabled,
            "status": project.status,
            "remarks": project.remarks,
            "created_at": project.created_at,
            **self._summary(rows),
        }

    def list_projects(self, ctx: TenantContext, *, company_account_id: UUID | None = None) -> list[dict[str, Any]]:
        cid = self._company(ctx)
        stmt = select(ProcDeliveryProject).where(
            ProcDeliveryProject.tenant_id == ctx.tenant_id,
            ProcDeliveryProject.company_id == cid,
            ProcDeliveryProject.is_deleted.is_(False),
        )
        if company_account_id is not None:
            stmt = stmt.where(ProcDeliveryProject.company_account_id == company_account_id)
        rows = self._db.scalars(stmt.order_by(ProcDeliveryProject.created_at.desc())).all()
        return [self._project_dto(row) for row in rows]

    def get_project(self, ctx: TenantContext, project_id: UUID) -> dict[str, Any]:
        project = self._get(ctx, project_id)
        sites = self._sites(project.id)
        return {**self._project_dto(project, sites), "sites": [self._site_dto(s) for s in sites]}

    def _next_code(self, company_id: UUID) -> str:
        prefix = f"DP-{date.today().year}-"
        count = self._db.scalar(
            select(func.count())
            .select_from(ProcDeliveryProject)
            .where(ProcDeliveryProject.company_id == company_id, ProcDeliveryProject.project_code.like(f"{prefix}%"))
        )
        return f"{prefix}{int(count or 0) + 1}"

    def create_project(
        self,
        ctx: TenantContext,
        *,
        name: str,
        opportunity_id: UUID | None = None,
        company_account_id: UUID | None = None,
        customer_name: str | None = None,
        remarks: str | None = None,
    ) -> dict[str, Any]:
        cid = self._company(ctx)
        title = (name or "").strip()
        if not title:
            raise ConflictException("Project name is required")
        if opportunity_id is not None:
            opp = self._crm.get_opportunity_brief(ctx, opportunity_id)
            if opp is None:
                raise NotFoundException("Opportunity not found")
            company_account_id = company_account_id or opp["company_account_id"]
            customer_name = customer_name or opp["customer_name"]
        from modules.procurement.service.service_contract_service import ServiceContractService

        project = ProcDeliveryProject(
            project_code=self._next_code(cid),
            name=title[:255],
            company_account_id=company_account_id,
            customer_name=(customer_name or "").strip() or None,
            opportunity_id=opportunity_id,
            tracking_token=secrets.token_urlsafe(24),
            public_tracking_enabled=True,
            status="active",
            remarks=(remarks or "").strip() or None,
            tenant_id=ctx.tenant_id,
            company_id=cid,
            branch_id=ServiceContractService(self._db)._branch(ctx, cid),
            created_by=ctx.user_id,
            updated_by=ctx.user_id,
        )
        self._db.add(project)
        self._db.flush()
        if opportunity_id is not None:
            self.import_sites_from_opportunity(ctx, project.id)
        return self.get_project(ctx, project.id)

    def update_project(self, ctx: TenantContext, project_id: UUID, **fields: Any) -> dict[str, Any]:
        project = self._get(ctx, project_id)
        if "status" in fields and fields["status"] not in PROJECT_STATUSES:
            raise ConflictException(f"status must be one of {', '.join(PROJECT_STATUSES)}")
        if fields.pop("rotate_tracking_token", False):
            project.tracking_token = secrets.token_urlsafe(24)
        for key in ("name", "customer_name", "remarks", "status", "public_tracking_enabled"):
            if key in fields and fields[key] is not None:
                setattr(project, key, fields[key])
        project.updated_by = ctx.user_id
        project.updated_at = _now()
        project.version = int(project.version or 1) + 1
        self._db.flush()
        return self.get_project(ctx, project_id)

    # -- sites --------------------------------------------------------------------
    def _site_dto(self, site: ProcDeliveryProjectSite) -> dict[str, Any]:
        today = date.today()
        return {
            "id": site.id,
            "project_id": site.project_id,
            "circle": site.circle,
            "site_code": site.site_code,
            "site_name": site.site_name,
            "address": site.address,
            "state": site.state,
            "gstin": site.gstin,
            "customer_po_number": site.customer_po_number,
            "ovf_id": site.ovf_id,
            "order_header_id": site.order_header_id,
            "item_summary": site.item_summary,
            "quantity": site.quantity,
            "milestone": site.milestone,
            "milestone_label": MILESTONE_LABELS.get(site.milestone, site.milestone),
            "status": site.status,
            "expected_delivery_date": site.expected_delivery_date,
            "actual_delivery_date": site.actual_delivery_date,
            "awb_number": site.awb_number,
            "delay_reason": site.delay_reason,
            "last_note": site.last_note,
            "delayed": bool(
                site.status not in ("delivered", "installed", "cancelled")
                and site.expected_delivery_date
                and site.expected_delivery_date < today
            ),
            "history": site.history or [],
            "updated_at": site.updated_at,
        }

    def _get_site(self, ctx: TenantContext, site_id: UUID) -> ProcDeliveryProjectSite:
        site = self._db.get(ProcDeliveryProjectSite, site_id)
        if site is None or site.is_deleted or site.tenant_id != ctx.tenant_id:
            raise NotFoundException("Site not found")
        self._get(ctx, site.project_id)
        return site

    def _apply(self, ctx: TenantContext, site: ProcDeliveryProjectSite, fields: dict[str, Any], *, source: str) -> bool:
        """Apply edits with the tracker's rules; returns True if anything changed."""
        changes: dict[str, Any] = {}
        for key, value in fields.items():
            if key not in _EDITABLE:
                continue
            if key not in _TYPED_FIELDS and value is not None and not isinstance(value, str):
                value = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
            if isinstance(value, str):
                value = value.strip() or None
            if key == "milestone":
                value = _milestone_from(value) if value else None
                if value is None:
                    continue
            elif key == "status" and value is not None and value not in SITE_STATUSES:
                raise ValidationException(f"Status must be one of {', '.join(SITE_STATUSES)}")
            elif key in ("expected_delivery_date", "actual_delivery_date"):
                value = _date_from(value)
            elif key == "quantity" and value is not None:
                try:
                    value = Decimal(str(value))
                except InvalidOperation as exc:
                    raise ValidationException(f"Quantity '{value}' is not a number") from exc
            elif key == "gstin" and value:
                value = str(value).upper()[:15]
            if key == "site_name" and not value:
                raise ValidationException("Site name cannot be empty")
            if getattr(site, key) != value:
                changes[key] = value

        if not changes:
            return False
        new_expected = changes.get("expected_delivery_date")
        if (
            new_expected
            and site.expected_delivery_date
            and new_expected > site.expected_delivery_date
            and not (changes.get("delay_reason") or fields.get("delay_reason"))
        ):
            raise ValidationException(
                f"{site.site_name}: expected delivery moved later - add a delay reason"
            )
        if "milestone" in changes:
            if "status" not in changes:
                changes["status"] = _status_for_milestone(changes["milestone"], site.status)
            if changes["milestone"] in ("reached_site", "installed") and not (
                changes.get("actual_delivery_date") or site.actual_delivery_date
            ):
                changes["actual_delivery_date"] = date.today()

        history = list(site.history or [])
        history.append(
            {
                "at": _now().isoformat(),
                "by": str(ctx.user_id) if ctx.user_id else None,
                "source": source,
                "changes": {
                    k: (v.isoformat() if isinstance(v, date) else str(v) if isinstance(v, Decimal) else v)
                    for k, v in changes.items()
                },
            }
        )
        for key, value in changes.items():
            setattr(site, key, value)
        site.history = history[-50:]
        site.updated_by = ctx.user_id
        site.updated_at = _now()
        site.version = int(site.version or 1) + 1
        return True

    def add_site(self, ctx: TenantContext, project_id: UUID, **fields: Any) -> dict[str, Any]:
        return self._site_dto(self._create_site(ctx, self._get(ctx, project_id), fields))

    def _create_site(
        self, ctx: TenantContext, project: ProcDeliveryProject, fields: dict[str, Any]
    ) -> ProcDeliveryProjectSite:
        name = str(fields.get("site_name") or "").strip()
        if not name:
            raise ValidationException("Site name is required")
        site = ProcDeliveryProjectSite(
            project_id=project.id,
            site_name=name[:255],
            milestone="order_placed",
            status="pending",
            sort_order=len(self._sites(project.id)) + 1,
            tenant_id=project.tenant_id,
            company_id=project.company_id,
            branch_id=project.branch_id,
            created_by=ctx.user_id,
            updated_by=ctx.user_id,
        )
        self._db.add(site)
        self._db.flush()
        self._apply(ctx, site, {k: v for k, v in fields.items() if k != "site_name"}, source="manual")
        for key in ("ovf_id", "order_header_id"):
            if fields.get(key):
                setattr(site, key, fields[key])
        self._db.flush()
        return site

    def update_site(self, ctx: TenantContext, site_id: UUID, **fields: Any) -> dict[str, Any]:
        site = self._get_site(ctx, site_id)
        self._apply(ctx, site, fields, source="manual")
        self._db.flush()
        return self._site_dto(site)

    def delete_site(self, ctx: TenantContext, site_id: UUID) -> None:
        site = self._get_site(ctx, site_id)
        site.is_deleted = True
        site.deleted_at = _now()
        site.deleted_by = ctx.user_id
        self._db.flush()

    def import_sites_from_opportunity(self, ctx: TenantContext, project_id: UUID) -> dict[str, int]:
        """One site per OVF on the project's deal (each OVF = one customer PO / location)."""
        project = self._get(ctx, project_id)
        if project.opportunity_id is None:
            raise ConflictException("Link the project to an opportunity first")
        existing = {s.ovf_id for s in self._sites(project.id) if s.ovf_id}
        created = 0
        for ovf in self._crm.list_ovf_briefs_for_opportunity(ctx, project.opportunity_id):
            if ovf["id"] in existing:
                continue
            order_id = self._db.scalar(
                select(ProcOrderHeader.id).where(
                    ProcOrderHeader.tenant_id == ctx.tenant_id,
                    ProcOrderHeader.is_deleted.is_(False),
                    ProcOrderHeader.source_module == "crm",
                    ProcOrderHeader.source_document_type == "ovf",
                    ProcOrderHeader.source_document_id == ovf["id"],
                )
            )
            self.add_site(
                ctx,
                project.id,
                site_name=ovf["shipping_address"] or ovf["shipping_state"] or ovf["ovf_no"],
                address=ovf["shipping_address"],
                state=ovf["shipping_state"],
                customer_po_number=ovf["po_number"],
                item_summary=ovf["item_summary"],
                quantity=ovf["quantity"],
                expected_delivery_date=ovf["expected_delivery_date"],
                ovf_id=ovf["id"],
                order_header_id=order_id,
            )
            created += 1
        return {"created": created}

    def apply_order_milestone(self, ctx: TenantContext, order: ProcOrderHeader) -> int:
        """Vendor PO milestone moved forward → sites on that PO follow."""
        if not order.delivery_milestone:
            return 0
        stmt = select(ProcDeliveryProjectSite).where(
            ProcDeliveryProjectSite.tenant_id == order.tenant_id,
            ProcDeliveryProjectSite.order_header_id == order.id,
            ProcDeliveryProjectSite.is_deleted.is_(False),
        )
        moved = 0
        target = DELIVERY_MILESTONES.index(order.delivery_milestone)
        for site in self._db.scalars(stmt).all():
            if DELIVERY_MILESTONES.index(site.milestone) >= target:
                continue
            fields: dict[str, Any] = {"milestone": order.delivery_milestone}
            if order.awb_number:
                fields["awb_number"] = order.awb_number
            if order.actual_delivery_date:
                fields["actual_delivery_date"] = order.actual_delivery_date
            if self._apply(ctx, site, fields, source="vendor_po"):
                moved += 1
        self._db.flush()
        return moved

    # -- Excel round-trip -----------------------------------------------------------
    def export_excel(self, ctx: TenantContext, project_id: UUID) -> tuple[str, bytes]:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation

        project = self._get(ctx, project_id)
        sites = self._sites(project.id)
        wb = Workbook()
        ws = wb.active
        ws.title = "Sites"
        ws.append([label for _key, label in EXCEL_COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E3A8A")
        for site in sites:
            dto = self._site_dto(site)
            row = []
            for key, _label in EXCEL_COLUMNS:
                if key == "site_key":
                    row.append(str(site.id))
                elif key == "milestone":
                    row.append(MILESTONE_LABELS.get(site.milestone, site.milestone))
                elif key == "quantity":
                    row.append(float(site.quantity) if site.quantity is not None else None)
                else:
                    value = dto.get(key)
                    row.append(value.strftime("%d-%m-%Y") if isinstance(value, date) else value)
            ws.append(row)
        milestone_col = [k for k, _ in EXCEL_COLUMNS].index("milestone") + 1
        status_col = [k for k, _ in EXCEL_COLUMNS].index("status") + 1
        last_row = max(len(sites) + 200, 200)
        milestones = DataValidation(type="list", formula1='"' + ",".join(MILESTONE_LABELS.values()) + '"')
        statuses = DataValidation(type="list", formula1='"' + ",".join(SITE_STATUSES) + '"')
        ws.add_data_validation(milestones)
        ws.add_data_validation(statuses)
        col = ws.cell(row=1, column=milestone_col).column_letter
        milestones.add(f"{col}2:{col}{last_row}")
        col = ws.cell(row=1, column=status_col).column_letter
        statuses.add(f"{col}2:{col}{last_row}")
        ws.column_dimensions["A"].hidden = True
        for column_cells in ws.columns:
            letter = column_cells[0].column_letter
            if letter != "A":
                ws.column_dimensions[letter].width = 18
        ws.freeze_panes = "B2"
        buf = BytesIO()
        wb.save(buf)
        return f"{project.project_code}-tracker.xlsx", buf.getvalue()

    def import_excel(self, ctx: TenantContext, project_id: UUID, *, content_base64: str) -> dict[str, Any]:
        from openpyxl import load_workbook

        project = self._get(ctx, project_id)
        try:
            raw = base64.b64decode(content_base64, validate=True)
        except (ValueError, binascii.Error) as exc:
            raise ValidationException("Invalid Excel file data") from exc
        if len(raw) > _MAX_IMPORT_BYTES:
            raise ValidationException("Tracker file is too large (max 8 MB)")
        try:
            ws = load_workbook(BytesIO(raw), data_only=True).worksheets[0]
        except Exception as exc:
            raise ValidationException("Could not read the Excel file - upload the downloaded .xlsx tracker") from exc

        header = [str(c.value or "").strip().lower() for c in ws[1]]
        by_label = {label.lower(): key for key, label in EXCEL_COLUMNS}
        by_label.update({key: key for key, _ in EXCEL_COLUMNS})
        columns = [by_label.get(h) for h in header]
        if "site_name" not in columns and "site_key" not in columns:
            raise ValidationException("The sheet needs a 'Site Key' or 'Site Name' column")

        sites = self._sites(project.id)
        by_id = {str(s.id): s for s in sites}
        by_po_code = {((s.customer_po_number or "").lower(), (s.site_code or "").lower()): s for s in sites}
        summary: dict[str, Any] = {"updated": 0, "created": 0, "unchanged": 0, "errors": []}
        for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            values = {col: row[i] for i, col in enumerate(columns) if col and i < len(row)}
            if not any(v not in (None, "") for v in values.values()):
                continue
            key = str(values.pop("site_key", "") or "").strip()
            site = by_id.get(key) if key else None
            if site is None:
                po = str(values.get("customer_po_number") or "").strip().lower()
                code = str(values.get("site_code") or "").strip().lower()
                if po or code:
                    site = by_po_code.get((po, code))
            try:
                with self._db.begin_nested():
                    if site is None:
                        if not str(values.get("site_name") or "").strip():
                            raise ValidationException("New rows need a Site Name")
                        site = self._create_site(ctx, project, values)
                        by_id[str(site.id)] = site
                        by_po_code[((site.customer_po_number or "").lower(), (site.site_code or "").lower())] = site
                        summary["created"] += 1
                    elif self._apply(ctx, site, values, source="excel"):
                        self._db.flush()
                        summary["updated"] += 1
                    else:
                        summary["unchanged"] += 1
            except (ValidationException, ConflictException) as exc:
                summary["errors"].append({"row": index, "message": str(getattr(exc, "message", None) or exc)})
        return summary

    # -- public view ----------------------------------------------------------------
    def public_view(self, token: str) -> dict[str, Any]:
        """Customer-facing tracker: milestones and dates only, no prices or vendors."""
        text = (token or "").strip()
        if len(text) < 20:
            raise NotFoundException("Tracker not found")
        project = self._db.scalar(
            select(ProcDeliveryProject).where(
                ProcDeliveryProject.tracking_token == text,
                ProcDeliveryProject.is_deleted.is_(False),
                ProcDeliveryProject.public_tracking_enabled.is_(True),
            )
        )
        if project is None:
            raise NotFoundException("Tracker not found")
        sites = [s for s in self._sites(project.id) if s.status != "cancelled"]
        today = date.today()
        return {
            "project_code": project.project_code,
            "name": project.name,
            "customer_name": project.customer_name,
            "status": project.status,
            **self._summary(sites),
            "sites": [
                {
                    "circle": s.circle,
                    "site_code": s.site_code,
                    "site_name": s.site_name,
                    "customer_po_number": s.customer_po_number,
                    "milestone": s.milestone,
                    "milestone_label": MILESTONE_LABELS.get(s.milestone, s.milestone),
                    "milestone_step": DELIVERY_MILESTONES.index(s.milestone) + 1,
                    "milestone_steps": len(DELIVERY_MILESTONES),
                    "status": s.status,
                    "expected_delivery_date": s.expected_delivery_date,
                    "actual_delivery_date": s.actual_delivery_date,
                    "awb_number": s.awb_number,
                    "delayed": bool(
                        s.status not in ("delivered", "installed")
                        and s.expected_delivery_date
                        and s.expected_delivery_date < today
                    ),
                    "updated_at": s.updated_at,
                }
                for s in sites
            ],
        }
