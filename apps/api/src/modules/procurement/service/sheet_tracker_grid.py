"""Parse Excel sheets and merge new columns/rows into an extracted tracker table."""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from core.exceptions import ValidationException

MAX_COLUMNS = 80
MAX_ROWS = 20_000
MAX_CELL = 4000

_SLUG_RE = re.compile(r"[^a-z0-9]+")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.hour == 0 and value.minute == 0 else value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        text = format(value, "f").rstrip("0").rstrip(".")
        return text or "0"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    text = str(value).strip()
    return text[:MAX_CELL]


def _slug(label: str, used: set[str]) -> str:
    base = _SLUG_RE.sub("_", label.strip().lower()).strip("_") or "col"
    cid = base[:60]
    n = 2
    while cid in used:
        cid = f"{base[:50]}_{n}"
        n += 1
    used.add(cid)
    return cid


def extract_sheet_table(raw: bytes) -> dict[str, Any]:
    """Read the first worksheet: header row + data rows."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationException("Excel support is not installed on the API") from exc

    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationException("Could not read the Excel file. Upload .xlsx or .xlsm.") from exc

    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise ValidationException("The workbook has no sheets")
        matrix = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    header_idx = None
    headers: list[str] = []
    for idx, row in enumerate(matrix[:40]):
        cells = [_cell_text(c) for c in (row or ())]
        filled = [c for c in cells if c]
        if len(filled) < 1:
            continue
        headers = cells
        header_idx = idx
        break
    if header_idx is None or not any(headers):
        raise ValidationException("No header row found on the first sheet")

    # Trim trailing empty header cells.
    last = 0
    for i, label in enumerate(headers):
        if label:
            last = i
    headers = headers[: last + 1]
    if len(headers) > MAX_COLUMNS:
        raise ValidationException(f"Trackers can have at most {MAX_COLUMNS} columns")

    used: set[str] = set()
    columns: list[dict[str, str]] = []
    seen_labels: dict[str, int] = {}
    for label in headers:
        display = label or f"Column {len(columns) + 1}"
        key = display.casefold()
        if key in seen_labels:
            display = f"{display} ({seen_labels[key] + 1})"
            seen_labels[key] += 1
        else:
            seen_labels[key] = 1
        columns.append({"id": _slug(display, used), "label": display[:180]})

    rows: list[dict[str, str]] = []
    for raw_row in matrix[header_idx + 1 :]:
        cells = [_cell_text(c) for c in (raw_row or ())]
        if not any(cells[: len(columns)]):
            continue
        item: dict[str, str] = {}
        for i, col in enumerate(columns):
            item[col["id"]] = cells[i] if i < len(cells) else ""
        rows.append(item)
        if len(rows) > MAX_ROWS:
            raise ValidationException(f"Trackers can have at most {MAX_ROWS} rows")

    if not rows:
        raise ValidationException("The sheet has headers but no data rows")

    return {"columns": columns, "rows": rows}


def _row_key(columns: list[dict[str, str]], row: dict[str, str]) -> str:
    if not columns:
        return ""
    first = (row.get(columns[0]["id"]) or "").strip().casefold()
    if first:
        return f"k:{first}"
    parts = [(row.get(c["id"]) or "").strip().casefold() for c in columns]
    return "f:" + "\x1f".join(parts)


def merge_tracker_table(
    existing_columns: list[dict[str, str]],
    existing_rows: list[dict[str, str]],
    incoming: dict[str, Any],
) -> dict[str, Any]:
    """Union columns by header label; append unseen rows; fill new cells on matches."""
    incoming_columns = list(incoming.get("columns") or [])
    incoming_rows = list(incoming.get("rows") or [])
    if not incoming_columns:
        raise ValidationException("Incoming sheet has no columns")

    columns = [{"id": str(c["id"]), "label": str(c["label"])} for c in existing_columns]
    used_ids = {c["id"] for c in columns}
    label_to_id = {c["label"].strip().casefold(): c["id"] for c in columns}

    added_column_labels: list[str] = []
    incoming_id_map: dict[str, str] = {}
    for col in incoming_columns:
        label = str(col.get("label") or "").strip() or "Column"
        src_id = str(col.get("id") or "")
        dest = label_to_id.get(label.casefold())
        if dest is None:
            dest = _slug(label, used_ids)
            columns.append({"id": dest, "label": label[:180]})
            label_to_id[label.casefold()] = dest
            added_column_labels.append(label)
        incoming_id_map[src_id] = dest

    if len(columns) > MAX_COLUMNS:
        raise ValidationException(f"Trackers can have at most {MAX_COLUMNS} columns")

    rows = [dict(r) for r in existing_rows]
    for row in rows:
        for col in columns:
            row.setdefault(col["id"], "")

    index = {_row_key(columns, row): i for i, row in enumerate(rows)}
    added_rows = 0
    updated_rows = 0

    for raw in incoming_rows:
        mapped: dict[str, str] = {c["id"]: "" for c in columns}
        for src_id, dest_id in incoming_id_map.items():
            mapped[dest_id] = str(raw.get(src_id) or "")
        if not any(mapped.values()):
            continue
        key = _row_key(columns, mapped)
        if key in index:
            target = rows[index[key]]
            changed = False
            for col in columns:
                incoming_val = mapped.get(col["id"], "")
                if not incoming_val:
                    continue
                current = target.get(col["id"], "")
                if current == incoming_val:
                    continue
                if current == "":
                    target[col["id"]] = incoming_val
                    changed = True
            if changed:
                updated_rows += 1
            continue
        if len(rows) >= MAX_ROWS:
            raise ValidationException(f"Trackers can have at most {MAX_ROWS} rows")
        rows.append(mapped)
        index[key] = len(rows) - 1
        added_rows += 1

    return {
        "columns": columns,
        "rows": rows,
        "added_columns": added_column_labels,
        "added_rows": added_rows,
        "updated_rows": updated_rows,
    }
