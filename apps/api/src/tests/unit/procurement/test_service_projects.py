"""Service rate contracts and multi-site delivery project tracker."""

import base64
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock
from uuid import uuid4

import pytest

from core.exceptions import ValidationException
from modules.procurement.service.delivery_project_service import (
    EXCEL_COLUMNS,
    DeliveryProjectService,
    _date_from,
    _milestone_from,
    _status_for_milestone,
)
from modules.procurement.service.service_contract_service import plan_total


def test_plan_total_prices_visits_plus_consumables():
    assert plan_total(12, Decimal("1500"), Decimal("2500")) == Decimal("20500.0000")
    assert plan_total(3, Decimal("999.50"), Decimal("0")) == Decimal("2998.5000")


def test_milestone_accepts_key_or_label():
    assert _milestone_from("dispatched_to_site") == "dispatched_to_site"
    assert _milestone_from("  Reached site ") == "reached_site"
    assert _milestone_from("") is None
    with pytest.raises(ValidationException):
        _milestone_from("somewhere")


def test_date_reads_common_formats():
    assert _date_from("2026-09-30") == date(2026, 9, 30)
    assert _date_from("30-09-2026") == date(2026, 9, 30)
    assert _date_from("30/09/2026") == date(2026, 9, 30)
    assert _date_from(datetime(2026, 9, 30, 10, 0)) == date(2026, 9, 30)
    assert _date_from(None) is None
    with pytest.raises(ValidationException):
        _date_from("next week")


def test_status_follows_milestone_unless_held():
    assert _status_for_milestone("order_placed", "pending") == "pending"
    assert _status_for_milestone("dispatched_to_site", "pending") == "in_progress"
    assert _status_for_milestone("reached_site", "in_progress") == "delivered"
    assert _status_for_milestone("installed", "delivered") == "installed"
    assert _status_for_milestone("installed", "on_hold") == "on_hold"


def _site(**overrides):
    base = dict(
        id=uuid4(),
        project_id=uuid4(),
        circle="North",
        site_code="1001",
        site_name="Gurgaon DC",
        address=None,
        state="Haryana",
        gstin=None,
        customer_po_number="PO-77",
        ovf_id=None,
        order_header_id=None,
        item_summary="Switch x 2",
        quantity=Decimal("2"),
        milestone="order_placed",
        status="pending",
        expected_delivery_date=date.today() + timedelta(days=10),
        actual_delivery_date=None,
        awb_number=None,
        delay_reason=None,
        last_note=None,
        history=[],
        updated_at=None,
        updated_by=None,
        version=1,
    )
    base.update(overrides)
    return SimpleNamespace(**base)


def _service(sites=None):
    svc = DeliveryProjectService.__new__(DeliveryProjectService)
    svc._db = MagicMock()
    project = SimpleNamespace(
        id=uuid4(), project_code="DP-2026-1", tenant_id=uuid4(), company_id=uuid4(), branch_id=uuid4()
    )
    svc._get = lambda ctx, project_id: project
    svc._sites = lambda project_id: list(sites or [])
    return svc


CTX = SimpleNamespace(user_id=uuid4(), tenant_id=uuid4())


def test_milestone_update_sets_status_actual_date_and_history():
    site = _site()
    assert _service()._apply(CTX, site, {"milestone": "Reached site"}, source="manual")
    assert site.milestone == "reached_site"
    assert site.status == "delivered"
    assert site.actual_delivery_date == date.today()
    assert site.history[-1]["changes"]["milestone"] == "reached_site"


def test_pushing_expected_date_later_needs_a_reason():
    site = _site()
    later = site.expected_delivery_date + timedelta(days=7)
    with pytest.raises(ValidationException):
        _service()._apply(CTX, site, {"expected_delivery_date": later}, source="manual")
    assert _service()._apply(
        CTX, site, {"expected_delivery_date": later, "delay_reason": "Customs hold"}, source="manual"
    )
    assert site.expected_delivery_date == later


def test_numeric_excel_cells_do_not_register_as_changes():
    site = _site()
    assert not _service()._apply(CTX, site, {"site_code": 1001.0, "quantity": 2}, source="excel")


def test_excel_round_trip_updates_by_site_key():
    from openpyxl import load_workbook

    site = _site()
    svc = _service([site])
    _name, content = svc.export_excel(CTX, uuid4())

    wb = load_workbook(BytesIO(content))
    ws = wb.active
    keys = [k for k, _ in EXCEL_COLUMNS]
    assert ws.cell(row=2, column=1).value == str(site.id)
    ws.cell(row=2, column=keys.index("milestone") + 1).value = "Dispatched to site"
    ws.cell(row=2, column=keys.index("awb_number") + 1).value = "AWB-123"
    buf = BytesIO()
    wb.save(buf)

    result = svc.import_excel(CTX, uuid4(), content_base64=base64.b64encode(buf.getvalue()).decode())
    assert result == {"updated": 1, "created": 0, "unchanged": 0, "errors": []}
    assert site.milestone == "dispatched_to_site"
    assert site.status == "in_progress"
    assert site.awb_number == "AWB-123"


def test_excel_row_errors_are_reported_per_row():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Site Name", "Milestone"])
    ws.append(["", "Installed"])
    ws.append(["Pune DC", "Teleported"])
    buf = BytesIO()
    wb.save(buf)

    result = _service([]).import_excel(CTX, uuid4(), content_base64=base64.b64encode(buf.getvalue()).decode())
    assert result["created"] == 0
    assert [e["row"] for e in result["errors"]] == [2, 3]
